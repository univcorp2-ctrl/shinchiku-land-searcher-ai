from __future__ import annotations

from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.cell import column_index_from_string, get_column_letter

from shinchiku_land_searcher.models import PropertyRecord


def read_property_records(
    path: Path,
    *,
    sheet_name: str | None = None,
    url_column: str = "A",
    header_row: int = 1,
    no_header: bool = False,
    require_url: bool = False,
) -> list[PropertyRecord]:
    """Read property rows from an Excel workbook.

    A common source file has URLs in column A and property attributes in later columns.
    The function keeps every cell value in ``raw`` so unknown columns are not lost.
    """

    if not path.exists():
        raise FileNotFoundError(f"Input Excel not found: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    worksheet = workbook[sheet_name] if sheet_name else workbook.active
    url_col_idx = column_index_from_string(url_column.upper())

    headers = _headers(worksheet, header_row=header_row, no_header=no_header)
    first_data_row = 1 if no_header else header_row + 1

    records: list[PropertyRecord] = []
    for row in range(first_data_row, worksheet.max_row + 1):
        raw: dict[str, Any] = {}
        row_has_value = False
        for col in range(1, worksheet.max_column + 1):
            value = worksheet.cell(row=row, column=col).value
            if value not in (None, ""):
                row_has_value = True
            key = headers.get(col, get_column_letter(col))
            raw[key] = value

        if not row_has_value:
            continue

        url_value = worksheet.cell(row=row, column=url_col_idx).value
        url = str(url_value or "").strip()
        if require_url and not is_probable_url(url):
            continue
        records.append(PropertyRecord(row_number=row, url=url, raw=raw))

    return records


def is_probable_url(value: str) -> bool:
    return value.startswith(("http://", "https://"))


def _headers(worksheet: Any, *, header_row: int, no_header: bool) -> dict[int, str]:
    headers: dict[int, str] = {}
    for col in range(1, worksheet.max_column + 1):
        if no_header:
            headers[col] = get_column_letter(col)
            continue
        value = worksheet.cell(row=header_row, column=col).value
        text = str(value).strip() if value not in (None, "") else get_column_letter(col)
        headers[col] = _dedupe_header(text, headers.values())
    return headers


def _dedupe_header(header: str, existing: Any) -> str:
    used = set(existing)
    if header not in used:
        return header
    i = 2
    while f"{header}_{i}" in used:
        i += 1
    return f"{header}_{i}"
